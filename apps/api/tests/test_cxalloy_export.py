"""The results export.

CxAlloy's API is read only, so this is how a ruling reaches the system of
record. The tests are mostly about the three properties ADR-0001 asks for:
determinism, failures kept separate, and delivery confirmed by a person.
"""

from __future__ import annotations

import io
import uuid
import zipfile
from datetime import UTC, datetime

import pytest
from sqlalchemy.orm import Session

from app.cxalloy import client as cx
from app.cxalloy.export import (
    FAILURES_NAME,
    MANIFEST_NAME,
    README_NAME,
    ExportError,
    build_export,
    confirm_delivery,
    delivery_status,
)
from app.models import (
    AppUser,
    Asset,
    CaptureRecipe,
    ChecklistItem,
    ExportStatus,
    Project,
    Ruling,
)
from app.models.enums import ChecklistItemState, CxAlloyDeliveryState, RulingVerdict, UserRole
from app.storage import InMemoryStorage
from tests import factories as f

BUCKET = "fie-exports"


@pytest.fixture
def storage() -> InMemoryStorage:
    return InMemoryStorage()


@pytest.fixture
def project(db: Session) -> Project:
    return f.make_project(db)


@pytest.fixture
def reviewer(db: Session) -> AppUser:
    return f.make_user(db, UserRole.REVIEWER)


def a_ruled_item(
    db: Session,
    project: Project,
    reviewer: AppUser,
    *,
    tag: str = "SWBD-101",
    state: ChecklistItemState = ChecklistItemState.REVIEWER_PASSED,
    note: str = "Label present and readable.",
) -> ChecklistItem:
    document = f.make_document(db, project)
    rule_set = f.make_rule_set(db, project, f"1.{uuid.uuid4().int % 1000}.0")
    requirement = f.make_requirement(db, project, document, rule_set=rule_set)
    asset = f.make_asset(db, project, tag=tag)
    asset.cxalloy_id = f"CX-{tag}"
    db.flush()

    item = f.make_checklist_item(
        db,
        asset,
        requirement,
        ruleset_version=rule_set.version,
        state=state,
        reviewer=reviewer.id,
        resolved_by=reviewer.id,
        resolved_at=datetime(2026, 9, 16, 9, 0, tzinfo=UTC),
    )
    db.add(
        Ruling(
            checklist_item_id=item.id,
            verdict=RulingVerdict.PASS
            if state is ChecklistItemState.REVIEWER_PASSED
            else RulingVerdict.FAIL,
            note=note,
            reviewer_id=reviewer.id,
        )
    )
    db.flush()
    return item


def read_zip(storage: InMemoryStorage, key: str) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(storage.get(BUCKET, key))) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


class TestTheReadClientCannotWrite:
    def test_no_client_exposes_a_write(self) -> None:
        """Read only by construction: not stubbed, absent."""
        forbidden = ("write", "create", "update", "post", "put", "delete", "attach", "issue")
        for klass in (cx.FileEquipmentSource, cx.UnconfiguredCxAlloyApi):
            names = [n for n in dir(klass) if not n.startswith("_")]
            assert not [n for n in names if any(w in n.lower() for w in forbidden)]

    def test_an_unconfigured_api_fails_loudly_rather_than_returning_nothing(self) -> None:
        """An empty equipment list looks exactly like a building with no equipment."""
        with pytest.raises(cx.CxAlloyUnavailableError, match="not configured"):
            cx.UnconfiguredCxAlloyApi().fetch_equipment("project")

    def test_an_equipment_export_can_be_read_whatever_the_headers_are_called(self) -> None:
        source = cx.FileEquipmentSource(
            b"Equipment Tag,CxAlloy ID,Equipment Class,System,Room\n"
            b"SWBD-101,CX-1,switchboard,normal power,Electrical Room 1-04\n"
        )
        (equipment,) = source.fetch_equipment("project")
        assert equipment.tag == "SWBD-101"
        assert equipment.cxalloy_id == "CX-1"
        assert equipment.room == "Electrical Room 1-04"

    def test_an_export_missing_the_tag_column_is_refused(self) -> None:
        with pytest.raises(cx.CxAlloyUnavailableError, match="tag"):
            cx.FileEquipmentSource(b"Something,Else\n1,2\n").fetch_equipment("project")


class TestThePackage:
    def test_a_ruling_becomes_a_manifest_row(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        a_ruled_item(db, project, reviewer)

        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)

        assert export.status is ExportStatus.RENDERED
        assert export.item_count == 1
        files = read_zip(storage, export.storage_key or "")
        manifest = files[MANIFEST_NAME].decode()
        assert "SWBD-101" in manifest
        assert "CX-SWBD-101" in manifest
        assert "reviewer_passed" in manifest

    def test_the_package_explains_itself(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        a_ruled_item(db, project, reviewer)
        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)

        readme = read_zip(storage, export.storage_key or "")[README_NAME].decode()
        assert "read only" in readme
        assert "confirm delivery" in readme

    def test_failures_get_their_own_file(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        """A failed item nobody has been told about is a defect, not admin lag."""
        a_ruled_item(db, project, reviewer, tag="SWBD-101")
        a_ruled_item(
            db,
            project,
            reviewer,
            tag="SWBD-102",
            state=ChecklistItemState.REVIEWER_FAILED,
            note="No label fitted.",
        )

        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)
        files = read_zip(storage, export.storage_key or "")

        assert export.failure_count == 1
        failures = files[FAILURES_NAME].decode()
        assert "SWBD-102" in failures
        assert "SWBD-101" not in failures
        assert "SWBD-101" in files[MANIFEST_NAME].decode(), "passes stay in the manifest"

    def test_the_failure_file_carries_why_it_matters(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        """Whoever opens the issue should not have to go and look it up."""
        a_ruled_item(db, project, reviewer, state=ChecklistItemState.REVIEWER_FAILED)
        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)

        failures = read_zip(storage, export.storage_key or "")[FAILURES_NAME].decode()
        assert "mislabelled panel" in failures


class TestDeterminism:
    def test_the_same_rulings_render_the_same_manifest(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        """A re-import has to be recognisable as a duplicate, not land twice."""
        a_ruled_item(db, project, reviewer, tag="SWBD-101")
        a_ruled_item(db, project, reviewer, tag="SWBD-102")

        first = build_export(db, storage, project_id=project.id, bucket=BUCKET)
        first_manifest = read_zip(storage, first.storage_key or "")[MANIFEST_NAME]

        # Put the items back as if the first export had never happened.
        for item in db.query(ChecklistItem).all():
            item.cxalloy_delivery_state = CxAlloyDeliveryState.NOT_APPLICABLE
            item.export_key = None
            item.cxalloy_exported_at = None
        db.flush()

        second = build_export(db, storage, project_id=project.id, bucket=BUCKET)
        second_manifest = read_zip(storage, second.storage_key or "")[MANIFEST_NAME]

        assert first_manifest == second_manifest
        assert first.content_hash == second.content_hash

    def test_the_export_key_is_stable_and_names_the_ruling(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        item = a_ruled_item(db, project, reviewer)
        ruling = db.query(Ruling).filter_by(checklist_item_id=item.id).one()

        build_export(db, storage, project_id=project.id, bucket=BUCKET)

        assert item.export_key == f"{item.id}:{ruling.id}"

    def test_a_correction_exports_as_a_new_row_not_a_replacement(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        """A correction is a new ruling; collapsing it onto the old one loses it."""
        item = a_ruled_item(db, project, reviewer)
        original = db.query(Ruling).filter_by(checklist_item_id=item.id).one()
        build_export(db, storage, project_id=project.id, bucket=BUCKET)
        first_key = item.export_key

        correction = Ruling(
            checklist_item_id=item.id,
            verdict=RulingVerdict.FAIL,
            note="On a second look the label is for the wrong voltage.",
            reviewer_id=reviewer.id,
            supersedes=original.id,
        )
        db.add(correction)
        item.state = ChecklistItemState.REVIEWER_FAILED
        item.cxalloy_delivery_state = CxAlloyDeliveryState.PENDING_EXPORT
        db.flush()

        build_export(db, storage, project_id=project.id, bucket=BUCKET)

        assert item.export_key != first_key


class TestNotExportingTheSameThingTwice:
    def test_an_exported_item_is_not_included_again(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        a_ruled_item(db, project, reviewer)

        first = build_export(db, storage, project_id=project.id, bucket=BUCKET)
        second = build_export(db, storage, project_id=project.id, bucket=BUCKET)

        assert first.item_count == 1
        assert second.item_count == 0

    def test_an_unruled_item_is_not_exported(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        document = f.make_document(db, project)
        rule_set = f.make_rule_set(db, project, "2.0.0")
        requirement = f.make_requirement(db, project, document, rule_set=rule_set)
        asset = f.make_asset(db, project, tag="SWBD-900")
        f.make_checklist_item(
            db,
            asset,
            requirement,
            ruleset_version=rule_set.version,
            state=ChecklistItemState.OPEN,
        )

        assert build_export(db, storage, project_id=project.id, bucket=BUCKET).item_count == 0


class TestDeliveryIsAPersonsAct:
    def test_rendering_a_file_is_not_delivery(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        item = a_ruled_item(db, project, reviewer)
        build_export(db, storage, project_id=project.id, bucket=BUCKET)

        assert item.cxalloy_delivery_state is CxAlloyDeliveryState.EXPORTED
        assert delivery_status(db, project.id).exported_not_confirmed == 1

    def test_a_person_confirming_clears_the_undelivered_count(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        item = a_ruled_item(db, project, reviewer)
        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)

        confirm_delivery(db, export.id, confirmed_by=reviewer.id)

        assert item.cxalloy_delivery_state is CxAlloyDeliveryState.DELIVERY_CONFIRMED
        assert item.cxalloy_delivery_confirmed_by == reviewer.id
        assert delivery_status(db, project.id).undelivered == 0

    def test_confirming_twice_is_refused(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        a_ruled_item(db, project, reviewer)
        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)
        confirm_delivery(db, export.id, confirmed_by=reviewer.id)

        with pytest.raises(ExportError, match="not rendered"):
            confirm_delivery(db, export.id, confirmed_by=reviewer.id)


class TestTheUndeliveredMetric:
    def test_failures_are_counted_separately_from_passes(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        """They alarm sooner: one is admin lag, the other is a defect nobody knows about."""
        a_ruled_item(db, project, reviewer, tag="SWBD-101")
        a_ruled_item(
            db, project, reviewer, tag="SWBD-102", state=ChecklistItemState.REVIEWER_FAILED
        )

        status = delivery_status(db, project.id)

        assert status.pending_export == 2
        assert status.undelivered_failures == 1

    def test_the_count_is_zero_when_everything_has_landed(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        a_ruled_item(db, project, reviewer)
        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)
        confirm_delivery(db, export.id, confirmed_by=reviewer.id)

        status = delivery_status(db, project.id)
        assert status.undelivered == 0
        assert status.undelivered_failures == 0


class TestMissingEvidence:
    def test_a_photo_that_never_uploaded_does_not_lose_the_ruling(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        """Q21: an event can reference a key with nothing behind it."""
        from app.capture.recipes import ensure_builtin_recipes

        item = a_ruled_item(db, project, reviewer)
        recipe: CaptureRecipe = ensure_builtin_recipes(db)[0]
        tech = f.make_user(db, UserRole.TECH)
        f.make_evidence(db, item, recipe, tech)  # storage has no bytes for it

        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)

        assert export.item_count == 1
        files = read_zip(storage, export.storage_key or "")
        assert MANIFEST_NAME in files
        assert not [n for n in files if n.startswith("photos/")]


class TestTheManualEntryWorklist:
    """Nothing can write to CxAlloy, so a person types these in. The worklist is
    the part they actually use."""

    def _worklist(self, storage: InMemoryStorage, key: str) -> object:
        from openpyxl import load_workbook

        from app.cxalloy.export import WORKLIST_NAME

        return load_workbook(io.BytesIO(read_zip(storage, key)[WORKLIST_NAME]))

    def test_the_package_leads_with_a_spreadsheet_not_a_csv(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        from app.cxalloy.export import WORKLIST_NAME

        a_ruled_item(db, project, reviewer)
        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)

        assert WORKLIST_NAME in read_zip(storage, export.storage_key or "")
        readme = read_zip(storage, export.storage_key or "")[README_NAME].decode()
        assert readme.index(WORKLIST_NAME) < readme.index(MANIFEST_NAME), "start here"

    def test_rows_are_ordered_by_cxalloy_id(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        """That is the order the person opens equipment in, so each is opened once."""
        for tag in ("SWBD-300", "SWBD-100", "SWBD-200"):
            a_ruled_item(db, project, reviewer, tag=tag)

        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)
        sheet = self._worklist(storage, export.storage_key or "")["Results"]  # type: ignore[index]

        ids = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
        assert ids == sorted(ids)

    def test_every_row_has_somewhere_to_keep_your_place(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        """Nobody gets through four hundred of these without being interrupted."""
        a_ruled_item(db, project, reviewer)
        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)
        sheet = self._worklist(storage, export.storage_key or "")["Results"]  # type: ignore[index]

        assert sheet.cell(row=1, column=1).value == "Done"
        assert sheet.cell(row=2, column=1).value in (None, "")
        assert sheet.freeze_panes == "A2", "the header stays visible while scrolling"

    def test_failures_are_on_their_own_sheet_and_marked_in_the_main_one(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        a_ruled_item(db, project, reviewer, tag="SWBD-101")
        a_ruled_item(
            db, project, reviewer, tag="SWBD-102", state=ChecklistItemState.REVIEWER_FAILED
        )

        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)
        workbook = self._worklist(storage, export.storage_key or "")

        failures = workbook["Failures"]  # type: ignore[index]
        assert failures.max_row == 2, "header plus the one failure"
        assert "SWBD-102" in [c.value for c in failures[2]]

        results = workbook["Results"]  # type: ignore[index]
        outcomes = {row[2]: row[4] for row in results.iter_rows(min_row=2, values_only=True)}
        assert outcomes["SWBD-102"] == "FAIL"
        assert outcomes["SWBD-101"] == "Pass"

    def test_the_failure_sheet_says_what_fails_if_it_is_missed(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        a_ruled_item(db, project, reviewer, state=ChecklistItemState.REVIEWER_FAILED)
        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)

        failures = self._worklist(storage, export.storage_key or "")["Failures"]  # type: ignore[index]
        assert any("mislabelled panel" in str(c.value or "") for c in failures[2])

    def test_an_asset_with_no_cxalloy_id_says_so_rather_than_showing_blank(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        """A blank cell reads as an oversight; the person needs to know it is not in there."""
        item = a_ruled_item(db, project, reviewer)
        asset = db.get(Asset, item.asset_id)
        assert asset is not None
        asset.cxalloy_id = None
        db.flush()

        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)
        sheet = self._worklist(storage, export.storage_key or "")["Results"]  # type: ignore[index]

        assert sheet.cell(row=2, column=2).value == "(not in CxAlloy)"

    def test_the_guide_says_the_platform_cannot_do_this_for_them(
        self, db: Session, storage: InMemoryStorage, project: Project, reviewer: AppUser
    ) -> None:
        a_ruled_item(db, project, reviewer)
        export = build_export(db, storage, project_id=project.id, bucket=BUCKET)
        guide = self._worklist(storage, export.storage_key or "")["How to use"]  # type: ignore[index]

        text = " ".join(str(row[0].value or "") for row in guide.iter_rows())
        assert "read only" in text
        assert "by hand" in text
        assert "confirm delivery" in text
