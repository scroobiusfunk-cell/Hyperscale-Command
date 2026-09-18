"""The manual-entry worklist.

Nothing can write to CxAlloy — not the API, not an importer. A person opens
CxAlloy and types the results in. That makes this package a worklist for a human
doing repetitive data entry, not a file for a machine, and the two want opposite
things: an importer wants one flat normalised table, a person wants the rows
grouped the way they will work through them, with somewhere to keep their place.

So the ordering is by CxAlloy id, because that is the order they will open
equipment in, and every row has a tick box. Someone doing four hundred of these
will be interrupted, and the cost of losing their place is doing part of it
twice.
"""

from __future__ import annotations

import io
from collections.abc import Callable
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

if TYPE_CHECKING:  # pragma: no cover
    from app.cxalloy.export import _Row
    from app.models import Evidence

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FAIL_FILL = PatternFill("solid", fgColor="FEE2E2")

RESULT_COLUMNS = (
    ("Done", 8),
    ("CxAlloy ID", 18),
    ("Asset tag", 16),
    ("Requirement", 60),
    ("Result", 16),
    ("Ruled on", 20),
    ("Reviewer note", 40),
    ("Evidence files", 44),
    ("Export key", 40),
)

FAILURE_COLUMNS = (
    ("Issue raised", 14),
    ("CxAlloy ID", 18),
    ("Asset tag", 16),
    ("Requirement", 50),
    ("What fails if this is missed", 50),
    ("Reviewer note", 40),
    ("Evidence files", 40),
    ("Export key", 40),
)

HOW_TO_USE = [
    "How to use this package",
    "",
    "Nothing here reaches CxAlloy on its own. This platform's access to CxAlloy is",
    "read only, so these results have to be entered by hand.",
    "",
    "1. Work down the Results sheet in CxAlloy ID order — that is the order you will",
    "   open the equipment in, so you only open each one once.",
    "2. Tick the Done column as you go. Nobody gets through four hundred of these",
    "   without being interrupted, and the tick is how you find your place again.",
    "3. The Failures sheet is the same items that failed, pulled out on their own.",
    "   Each one needs an issue raising against it. They are also in Results.",
    "4. The photos are in the photos/ folder, named by asset tag and checklist item.",
    "   The Evidence files column tells you which belong to which row.",
    "",
    "When you have finished, confirm delivery in the reviewer console. Until somebody",
    "does that, these rulings count as not yet in CxAlloy, and the undelivered count",
    "keeps climbing whether or not the work was actually done.",
    "",
    "Export key is a stable id for each ruling. If this package is produced twice, the",
    "keys are the same, so you can tell a repeat from something new.",
]


def _write_header(sheet: Worksheet, columns: tuple[tuple[str, int], ...]) -> None:
    for index, (title, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"


def build_workbook(rows: list[_Row], photo_path: Callable[[_Row, Evidence], str]) -> bytes:
    """Render the worklist. `photo_path` names the file for one piece of evidence."""
    workbook = Workbook()
    # openpyxl hands you a default sheet; removing it and creating all three
    # explicitly keeps the tab order deliberate and the types honest.
    workbook.remove(workbook.worksheets[0])

    guide = workbook.create_sheet("How to use")
    guide.column_dimensions["A"].width = 90
    for index, line in enumerate(HOW_TO_USE, start=1):
        cell = guide.cell(row=index, column=1, value=line)
        if index == 1:
            cell.font = Font(bold=True, size=14)

    results = workbook.create_sheet("Results")
    _write_header(results, RESULT_COLUMNS)

    failures = workbook.create_sheet("Failures")
    _write_header(failures, FAILURE_COLUMNS)

    # CxAlloy id order: the order the person will open equipment in.
    ordered = sorted(
        rows, key=lambda r: (r.asset.cxalloy_id or "~", r.asset.tag, r.requirement.statement)
    )

    for row in ordered:
        files = "\n".join(photo_path(row, e) for e in row.evidence)
        ruled_at = row.item.resolved_at.strftime("%Y-%m-%d %H:%M") if row.item.resolved_at else ""
        note = (row.ruling.note if row.ruling and row.ruling.note else "") or ""

        results.append(
            [
                "",
                row.asset.cxalloy_id or "(not in CxAlloy)",
                row.asset.tag,
                row.requirement.statement,
                "FAIL" if row.failed else "Pass",
                ruled_at,
                note,
                files,
                row.export_key,
            ]
        )
        if row.failed:
            for cell in results[results.max_row]:
                cell.fill = FAIL_FILL
            failures.append(
                [
                    "",
                    row.asset.cxalloy_id or "(not in CxAlloy)",
                    row.asset.tag,
                    row.requirement.statement,
                    row.requirement.why_it_matters,
                    note,
                    files,
                    row.export_key,
                ]
            )

    for sheet in (results, failures):
        for excel_row in sheet.iter_rows(min_row=2):
            for cell in excel_row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
